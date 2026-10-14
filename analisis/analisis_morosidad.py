#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Análisis estadístico de morosidad para escuelas primarias.

Bloques: carga y validación de datos, estadística descriptiva, estadística inferencial
(a nivel familia), validación del puntaje de riesgo (ROC), análisis de Pareto y
conclusiones automáticas en Markdown.

Uso:
    python analisis_morosidad.py                                   # busca datos/alumnos.csv, familias.csv y cuotas.csv; si no están, usa datos sintéticos
    python analisis_morosidad.py --alumnos a.csv --familias f.csv --cuotas c.csv
    python analisis_morosidad.py --archivo datos.xlsx              # Excel con hojas alumnos/familias/cuotas, o una sola tabla
    python analisis_morosidad.py --sintetico                       # fuerza el dataset sintético

Es la misma lógica que la pantalla "Estadística → Análisis de morosidad" del sistema
(src/morosidad.js); el botón "Exportar dataset para Python" baja los tres CSV que lee este script.
"""

from __future__ import annotations

import argparse
import logging
import math
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats
from statsmodels.stats.multitest import multipletests

import matplotlib

matplotlib.use("Agg")  # sin ventana: sólo se guardan los PNG
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.ticker import FuncFormatter, PercentFormatter  # noqa: E402

# =============================================================================
# CONFIGURACIÓN: se adapta acá, sin tocar el resto del código
# =============================================================================
CONFIG = {
    # Archivos por defecto (relativos a la carpeta desde donde se ejecuta)
    "ARCHIVOS": {"alumnos": "datos/alumnos.csv", "familias": "datos/familias.csv", "cuotas": "datos/cuotas.csv"},
    # Si los datos vienen en un Excel, nombre de cada hoja
    "HOJAS": {"alumnos": "alumnos", "familias": "familias", "cuotas": "cuotas"},
    # Campo lógico -> nombre real de la columna en tus archivos
    "COLUMNAS": {
        "alumnos": {
            "id_alumno": "id_alumno",
            "id_familia": "id_familia",
            "edad": "edad",
            "beca_pct": "beca_pct",
            "asistencia_30d_pct": "asistencia_30d_pct",
        },
        "familias": {
            "id_familia": "id_familia",
            "hijos_por_familia": "hijos_por_familia",
            "cuota_mensual_familiar": "cuota_mensual_familiar",
            "puntaje_riesgo_actual": "puntaje_riesgo_actual",
        },
        "cuotas": {
            "id_cuota": "id_cuota",
            "id_familia": "id_familia",
            "dias_atraso": "dias_atraso",
            "monto_adeudado": "monto_adeudado",  # opcional: si no existe, se omite el Pareto de deuda
        },
    },
    # Parámetros del análisis
    "UMBRAL_MOROSO_DIAS": 5,          # familia morosa: al menos una cuota con atraso mayor a este valor
    "ALFA": 0.05,
    "CORTES_RIESGO": (0.25, 0.75),    # percentiles del puntaje: bajo <= P25 < medio <= P75 < alto
    "CORTES_ASISTENCIA": (90, 95),    # % de asistencia promedio de la familia
    "PARETO_REFERENCIA": 0.80,
    "TOP_PUNTAJE": 0.20,              # fracción de familias con mayor puntaje para el Pareto del puntaje
    "BOOTSTRAP": 1000,
    "MONTE_CARLO": 10000,
    "SEMILLA": 20261014,
    # Cortes del semáforo que usa hoy el sistema (para comparar con los percentiles reales). None si no aplica.
    "UMBRALES_SEMAFORO_ACTUALES": {"medio": 25, "alto": 50},
    "SALIDA": "resultados",
}

log = logging.getLogger("morosidad")
ORDEN_RIESGO = ["Bajo", "Medio", "Alto"]
MOROSO = ["Morosa", "No morosa"]


# =============================================================================
# Formato rioplatense
# =============================================================================
def fnum(v, d: int = 1) -> str:
    """Número con coma decimal y punto de miles; sin ceros de relleno."""
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return "—"
    s = f"{float(v):,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if "," in s:
        s = s.rstrip("0").rstrip(",")
    return s


def fpct(v, d: int = 1) -> str:
    return "—" if v is None else f"{fnum(v * 100, d)}%"


def fmoney(v) -> str:
    return "—" if v is None or pd.isna(v) else "$" + fnum(round(v), 0)


def fp(nombre: str, p) -> str:
    if p is None or pd.isna(p):
        return f"{nombre} —"
    return f"{nombre} < 0,001" if p < 0.001 else f"{nombre} = {fnum(p, 3)}"


# =============================================================================
# 0. Carga de datos
# =============================================================================
def generar_sintetico(seed: int) -> dict[str, pd.DataFrame]:
    """Dataset reproducible con ~64 familias, ~94 alumnos y ~680 cuotas: muchos ceros en el atraso y cola larga."""
    rng = np.random.default_rng(seed)
    n_fam = 64
    hijos = rng.choice([1, 2, 3, 4], size=n_fam, p=[0.58, 0.31, 0.09, 0.02])
    # Propensión latente a atrasarse (lognormal: la mayoría paga bien, unas pocas familias concentran el atraso)
    latente = rng.lognormal(mean=0.0, sigma=1.0, size=n_fam)
    alumnos, familias, cuotas = [], [], []
    id_alu, id_cuota = 1, 1
    for f in range(n_fam):
        id_fam = f + 1
        becas = []
        for _ in range(hijos[f]):
            beca = rng.choice([0, 25, 50, 100], p=[0.94, 0.03, 0.02, 0.01])
            becas.append(beca)
            asis = float(np.clip(100 * rng.beta(28 - min(latente[f], 6), 1.6), 55, 100))
            alumnos.append({"id_alumno": id_alu, "id_familia": id_fam, "edad": int(rng.integers(6, 13)),
                            "beca_pct": beca, "asistencia_30d_pct": round(asis, 1)})
            id_alu += 1
        # Cuota con descuento por hermano (10% desde el segundo) y beca
        base = 185000
        cuota = sum(base * (1 - (0.10 if i else 0)) * (1 - b / 100) for i, b in enumerate(becas))
        puntaje = float(np.clip(round(12 * latente[f] + rng.normal(0, 6)), 0, 100))
        familias.append({"id_familia": id_fam, "hijos_por_familia": int(hijos[f]),
                         "cuota_mensual_familiar": round(cuota), "puntaje_riesgo_actual": puntaje})
        p_atraso = 1 - math.exp(-0.12 * latente[f] ** 1.5)  # probabilidad de pagar tarde una cuota (muchos ceros)
        for _ in range(int(rng.integers(10, 12))):
            if rng.random() < p_atraso:
                atraso = int(np.ceil(rng.lognormal(mean=1.2 + 0.25 * latente[f], sigma=1.1)))
            else:
                atraso = 0
            adeudado = round(cuota) if atraso > 15 and rng.random() < 0.5 else 0
            cuotas.append({"id_cuota": id_cuota, "id_familia": id_fam, "dias_atraso": atraso, "monto_adeudado": adeudado})
            id_cuota += 1
    return {"alumnos": pd.DataFrame(alumnos), "familias": pd.DataFrame(familias), "cuotas": pd.DataFrame(cuotas)}


def _leer(ruta: Path, hoja: str | None = None) -> pd.DataFrame:
    if ruta.suffix.lower() in (".xlsx", ".xls"):
        return pd.read_excel(ruta, sheet_name=hoja or 0)
    # sep=None detecta coma o punto y coma; utf-8-sig tolera el BOM que agrega Excel
    return pd.read_csv(ruta, sep=None, engine="python", encoding="utf-8-sig")


def _renombrar(df: pd.DataFrame, tabla: str) -> pd.DataFrame:
    """Pasa de los nombres reales (CONFIG['COLUMNAS']) a los campos lógicos."""
    mapa = CONFIG["COLUMNAS"][tabla]
    inverso = {real: logico for logico, real in mapa.items() if real in df.columns}
    faltan = [logico for logico, real in mapa.items() if real not in df.columns]
    if faltan:
        log.warning("Tabla %s: faltan las columnas %s (se tratarán como ausentes)", tabla, ", ".join(faltan))
    return df.rename(columns=inverso)[[c for c in mapa if c in inverso.values()]]


def _separar_tabla_unica(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """
    Un solo archivo plano: se reconstruyen las tres tablas por id_familia.
    Decisión documentada: cada fila con id_cuota es una cuota; cada id_alumno distinto es un alumno;
    los campos de familia se toman de la primera fila de cada id_familia (deberían repetirse iguales).
    """
    todas = {**CONFIG["COLUMNAS"]["alumnos"], **CONFIG["COLUMNAS"]["familias"], **CONFIG["COLUMNAS"]["cuotas"]}
    df = df.rename(columns={real: logico for logico, real in todas.items() if real in df.columns})
    out = {}
    cols_fam = [c for c in CONFIG["COLUMNAS"]["familias"] if c in df.columns]
    out["familias"] = df[cols_fam].drop_duplicates("id_familia", keep="first")
    if "id_alumno" in df.columns:
        cols = [c for c in CONFIG["COLUMNAS"]["alumnos"] if c in df.columns]
        out["alumnos"] = df[cols].dropna(subset=["id_alumno"]).drop_duplicates("id_alumno")
    else:
        out["alumnos"] = pd.DataFrame(columns=list(CONFIG["COLUMNAS"]["alumnos"]))
    if "id_cuota" in df.columns:
        cols = [c for c in CONFIG["COLUMNAS"]["cuotas"] if c in df.columns]
        out["cuotas"] = df[cols].dropna(subset=["id_cuota"]).drop_duplicates("id_cuota")
    else:
        out["cuotas"] = pd.DataFrame(columns=list(CONFIG["COLUMNAS"]["cuotas"]))
    log.info("Archivo único: se separó en %d familias, %d alumnos y %d cuotas (join por id_familia)",
             len(out["familias"]), len(out["alumnos"]), len(out["cuotas"]))
    return out


def cargar_datos(args) -> tuple[dict[str, pd.DataFrame], str]:
    """Lee CSV o Excel (tres tablas o una sola). Sin archivos, genera el dataset sintético."""
    if args.sintetico:
        log.info("Usando dataset sintético (semilla %s)", CONFIG["SEMILLA"])
        return generar_sintetico(CONFIG["SEMILLA"]), "sintético"
    if args.archivo:
        ruta = Path(args.archivo)
        if ruta.suffix.lower() in (".xlsx", ".xls"):
            hojas = pd.ExcelFile(ruta).sheet_names
            if all(CONFIG["HOJAS"][t] in hojas for t in ("alumnos", "familias", "cuotas")):
                datos = {t: _renombrar(_leer(ruta, CONFIG["HOJAS"][t]), t) for t in ("alumnos", "familias", "cuotas")}
                log.info("Excel con tres hojas: %s", ", ".join(CONFIG["HOJAS"].values()))
                return datos, str(ruta)
        return _separar_tabla_unica(_leer(ruta)), str(ruta)
    rutas = {t: Path(getattr(args, t) or CONFIG["ARCHIVOS"][t]) for t in ("alumnos", "familias", "cuotas")}
    if not rutas["familias"].exists() or not rutas["cuotas"].exists():
        log.warning("No se encontraron %s y %s: se usa el dataset sintético", rutas["familias"], rutas["cuotas"])
        return generar_sintetico(CONFIG["SEMILLA"]), "sintético"
    datos = {}
    for t, r in rutas.items():
        if r.exists():
            datos[t] = _renombrar(_leer(r), t)
        else:
            log.warning("No se encontró %s: la tabla %s queda vacía", r, t)
            datos[t] = pd.DataFrame(columns=list(CONFIG["COLUMNAS"][t]))
        log.info("Leído %s: %d filas", r, len(datos[t]))
    return datos, ", ".join(str(r) for r in rutas.values())


# =============================================================================
# 1. Validación
# =============================================================================
def validar(datos: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Nulos, duplicados, rangos imposibles e identificadores huérfanos. Informa; no aborta."""
    filas = []

    def add(tabla, chequeo, n, detalle=""):
        filas.append({"tabla": tabla, "chequeo": chequeo, "casos": int(n), "ok": n == 0, "detalle": detalle})

    def num(df, c):
        return pd.to_numeric(df[c], errors="coerce") if c in df.columns else pd.Series(dtype=float)

    al, fa, cu = datos["alumnos"], datos["familias"], datos["cuotas"]
    for tabla, df in (("alumnos", al), ("familias", fa), ("cuotas", cu)):
        for c in df.columns:
            if c == "beca_pct":
                continue  # vacío = sin beca
            add(tabla, f"Nulos en {c}", df[c].isna().sum())
    for tabla, df, k in (("alumnos", al, "id_alumno"), ("familias", fa, "id_familia"), ("cuotas", cu, "id_cuota")):
        if k in df.columns:
            add(tabla, f"Duplicados de {k}", df[k].dropna().duplicated().sum())

    def rango(tabla, df, c, lo, hi):
        if c in df.columns:
            v = num(df, c).dropna()
            add(tabla, f"{c} fuera de {lo} a {hi}", ((v < lo) | (v > hi)).sum())

    rango("alumnos", al, "asistencia_30d_pct", 0, 100)
    rango("alumnos", al, "beca_pct", 0, 100)
    rango("alumnos", al, "edad", 3, 18)
    rango("familias", fa, "puntaje_riesgo_actual", 0, 100)
    if "cuota_mensual_familiar" in fa.columns:
        add("familias", "Cuota mensual menor o igual a 0", (num(fa, "cuota_mensual_familiar") <= 0).sum(),
            "familias sin alumnos activos o con beca completa")
    if "dias_atraso" in cu.columns:
        add("cuotas", "Atraso negativo", (num(cu, "dias_atraso") < 0).sum())
    if "monto_adeudado" in cu.columns:
        add("cuotas", "Monto adeudado negativo", (num(cu, "monto_adeudado") < 0).sum())
    ids = set(fa["id_familia"]) if "id_familia" in fa.columns else set()
    if "id_familia" in cu.columns:
        add("cuotas", "Cuotas de familias inexistentes", (~cu["id_familia"].isin(ids)).sum())
    if "id_familia" in al.columns:
        add("alumnos", "Alumnos de familias inexistentes", (~al["id_familia"].isin(ids)).sum())
    out = pd.DataFrame(filas)
    for _, r in out[~out["ok"]].iterrows():
        log.warning("Validación · %s: %s → %d", r["tabla"], r["chequeo"], r["casos"])
    if out["ok"].all():
        log.info("Validación: sin observaciones")
    return out


def limpiar(datos: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    """Convierte a número y deja como faltantes los valores imposibles (no se descartan filas enteras)."""
    al, fa, cu = (datos[t].copy() for t in ("alumnos", "familias", "cuotas"))
    for df, cols in ((al, ["edad", "beca_pct", "asistencia_30d_pct"]),
                     (fa, ["hijos_por_familia", "cuota_mensual_familiar", "puntaje_riesgo_actual"]),
                     (cu, ["dias_atraso", "monto_adeudado"])):
        for c in cols:
            if c in df.columns:
                df[c] = pd.to_numeric(df[c], errors="coerce")
    if "beca_pct" in al.columns:
        al["beca_pct"] = al["beca_pct"].fillna(0)
    if "asistencia_30d_pct" in al.columns:
        al.loc[~al["asistencia_30d_pct"].between(0, 100), "asistencia_30d_pct"] = np.nan
    if "dias_atraso" in cu.columns:
        cu.loc[cu["dias_atraso"] < 0, "dias_atraso"] = np.nan
    if "monto_adeudado" in cu.columns:
        cu.loc[cu["monto_adeudado"] < 0, "monto_adeudado"] = np.nan
    return {"alumnos": al, "familias": fa.drop_duplicates("id_familia"), "cuotas": cu.drop_duplicates("id_cuota")}


def tabla_familias(datos: dict[str, pd.DataFrame]) -> tuple[pd.DataFrame, dict]:
    """Une alumnos y cuotas a la tabla de familias por id_familia y arma las variables del análisis."""
    al, fa, cu = datos["alumnos"], datos["familias"].copy(), datos["cuotas"]
    umbral = CONFIG["UMBRAL_MOROSO_DIAS"]
    hay_monto = "monto_adeudado" in cu.columns
    cu_ok = cu.dropna(subset=["dias_atraso"])
    agg = cu_ok.groupby("id_familia").agg(
        n_cuotas=("dias_atraso", "size"),
        atraso_max=("dias_atraso", "max"),
        atraso_acum=("dias_atraso", "sum"),
        moroso=("dias_atraso", lambda s: bool((s > umbral).any())),
        **({"monto_adeudado": ("monto_adeudado", "sum")} if hay_monto else {}),
    )
    if len(al) and "id_familia" in al.columns:
        a = al.groupby("id_familia").agg(
            becados=("beca_pct", lambda s: int((s > 0).sum())) if "beca_pct" in al.columns else ("id_familia", lambda s: 0),
            asistencia_prom=("asistencia_30d_pct", "mean") if "asistencia_30d_pct" in al.columns else ("id_familia", lambda s: np.nan),
        )
    else:
        a = pd.DataFrame(columns=["becados", "asistencia_prom"])
    F = fa.merge(agg, on="id_familia", how="left").merge(a, on="id_familia", how="left")
    F["n_cuotas"] = F["n_cuotas"].fillna(0).astype(int)
    F["becados"] = F["becados"].fillna(0).astype(int)
    F["beca"] = F["becados"] > 0
    F["beca_lbl"] = np.where(F["beca"], "Con beca", "Sin beca")
    total = len(F)
    # Sólo familias con al menos una cuota vencida y puntaje calculado
    F = F[(F["n_cuotas"] > 0) & F["puntaje_riesgo_actual"].notna()].copy()
    F["moroso"] = F["moroso"].astype(bool)
    F["moroso_lbl"] = np.where(F["moroso"], "Morosa", "No morosa")
    p25, p75 = (float(np.quantile(F["puntaje_riesgo_actual"], q)) for q in CONFIG["CORTES_RIESGO"])
    s = F["puntaje_riesgo_actual"]
    if p25 == p75:
        F["riesgo_cat"] = np.where(s <= p25, "Bajo", "Alto")
    else:
        F["riesgo_cat"] = np.select([s <= p25, s <= p75], ["Bajo", "Medio"], "Alto")
    ca, cb = CONFIG["CORTES_ASISTENCIA"]
    F["asistencia_cat"] = pd.Series(pd.NA, index=F.index, dtype=object)
    m = F["asistencia_prom"].notna()
    F.loc[m, "asistencia_cat"] = np.select([F.loc[m, "asistencia_prom"] < ca, F.loc[m, "asistencia_prom"] < cb],
                                           [f"Menos de {ca}%", f"{ca} a {cb}%"], f"{cb}% o más")
    # Hijos: 1, 2, 3 o más; si el grupo de arriba queda con menos de 3 familias, se une al anterior
    tope = 3
    while tope > 2 and (F["hijos_por_familia"] >= tope).sum() < 3:
        tope -= 1
    F["hijos_grupo"] = np.where(F["hijos_por_familia"] >= tope, f"{tope} o más", F["hijos_por_familia"].astype(int).astype(str))
    meta = {"p25": p25, "p75": p75, "hijos_tope": tope, "familias_total": total, "hay_monto": hay_monto}
    log.info("Tabla de familias: %d de %d con cuotas vencidas; %d morosas (atraso > %d días)",
             len(F), total, int(F["moroso"].sum()), umbral)
    return F, meta


# =============================================================================
# 2. Estadística descriptiva
# =============================================================================
def _describir(v: pd.Series) -> dict:
    v = pd.to_numeric(v, errors="coerce").dropna().astype(float)
    n = len(v)
    if n == 0:
        return {"N": 0}
    cuenta = v.value_counts()
    moda = float(cuenta[cuenta == cuenta.max()].index.min()) if cuenta.max() > 1 else np.nan
    media, desvio = v.mean(), v.std(ddof=1) if n > 1 else 0.0
    return {
        "N": n, "media": media, "mediana": v.median(), "moda": moda, "desvio": desvio,
        "CV_pct": desvio / abs(media) * 100 if media else np.nan,
        "min": v.min(), "P25": v.quantile(0.25), "P75": v.quantile(0.75), "max": v.max(),
        "asimetria": float(stats.skew(v, bias=False)) if n > 2 and desvio > 0 else np.nan,
    }


def descriptiva(datos: dict[str, pd.DataFrame], salida: Path, generados: list, omitidos: list) -> tuple[pd.DataFrame, dict]:
    """Tabla resumen por variable numérica, cada una con su N y su unidad de análisis."""
    al, fa, cu = datos["alumnos"], datos["familias"], datos["cuotas"]
    umbral = CONFIG["UMBRAL_MOROSO_DIAS"]
    defs = [
        ("Edad", "años", "alumnos", al, "edad", None),
        ("Beca (sólo alumnos becados)", "%", "alumnos", al, "beca_pct", lambda s: s > 0),
        ("Asistencia últimos 30 días", "%", "alumnos", al, "asistencia_30d_pct", None),
        ("Hijos por familia", "alumnos", "familias", fa, "hijos_por_familia", None),
        ("Cuota mensual familiar", "$", "familias", fa, "cuota_mensual_familiar", lambda s: s > 0),
        ("Puntaje de riesgo actual", "puntos", "familias", fa, "puntaje_riesgo_actual", None),
        ("Días de atraso por cuota", "días", "cuotas", cu, "dias_atraso", None),
        ("Monto adeudado por cuota", "$", "cuotas", cu, "monto_adeudado", None),
    ]
    filas = []
    for variable, unidad, ua, df, col, filtro in defs:
        if col not in df.columns:
            omitidos.append(f"Descriptiva de '{variable}': no existe la columna {col}")
            continue
        s = pd.to_numeric(df[col], errors="coerce")
        if filtro is not None:
            s = s[filtro(s)]
        d = _describir(s)
        asim = d.get("asimetria", np.nan)
        marcada = bool(pd.notna(asim) and abs(asim) > 1)
        filas.append({"variable": variable, "unidad": unidad, "unidad_analisis": ua, **d,
                      "asimetria_marcada": "sí" if marcada else "no",
                      "recomendacion": "Usar mediana y P25–P75 (la media se deforma por la cola)" if marcada else ""})
    desc = pd.DataFrame(filas)
    at = pd.to_numeric(cu["dias_atraso"], errors="coerce").dropna() if "dias_atraso" in cu.columns else pd.Series(dtype=float)
    extra = {
        "n": len(at),
        "pct_mayor_0": float((at > 0).mean()) if len(at) else None,
        "pct_mayor_umbral": float((at > umbral).mean()) if len(at) else None,
    }
    for _, r in desc[desc["asimetria_marcada"] == "sí"].iterrows():
        log.info("Asimetría marcada en %s (%.2f): se recomienda mediana y percentiles", r["variable"], r["asimetria"])
    log.info("Cuotas con atraso > 0: %s · con atraso > %d días: %s", fpct(extra["pct_mayor_0"]), umbral, fpct(extra["pct_mayor_umbral"]))

    # CSV para Excel en español: separador ; y coma decimal
    desc.to_csv(salida / "descriptiva.csv", sep=";", decimal=",", index=False, encoding="utf-8-sig", float_format="%.2f")
    generados.append("descriptiva.csv")
    _excel_descriptiva(desc, extra, salida / "descriptiva.xlsx", generados, omitidos)
    return desc, extra


def _excel_descriptiva(desc: pd.DataFrame, extra: dict, ruta: Path, generados: list, omitidos: list) -> None:
    """Excel legible: $ con separador de miles en montos, dos decimales en el resto (la coma la pone la configuración regional)."""
    try:
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        omitidos.append("descriptiva.xlsx: falta el paquete openpyxl (pip install openpyxl)")
        return
    nombres = {"variable": "Variable", "unidad": "Unidad", "unidad_analisis": "Unidad de análisis", "N": "N", "media": "Media",
               "mediana": "Mediana", "moda": "Moda", "desvio": "Desvío", "CV_pct": "CV (%)", "min": "Mín.", "P25": "P25",
               "P75": "P75", "max": "Máx.", "asimetria": "Asimetría", "asimetria_marcada": "|Asimetría| > 1",
               "recomendacion": "Recomendación"}
    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        desc.rename(columns=nombres).to_excel(xw, sheet_name="Descriptiva", index=False)
        pd.DataFrame([
            {"Indicador": "Cuotas analizadas", "Valor": extra["n"]},
            {"Indicador": "% de cuotas con atraso mayor a 0", "Valor": extra["pct_mayor_0"]},
            {"Indicador": f"% de cuotas con atraso mayor a {CONFIG['UMBRAL_MOROSO_DIAS']} días", "Valor": extra["pct_mayor_umbral"]},
        ]).to_excel(xw, sheet_name="Atraso", index=False)
        ws = xw.sheets["Descriptiva"]
        cols = list(nombres.values())
        num_cols = ["Media", "Mediana", "Moda", "Desvío", "Mín.", "P25", "P75", "Máx."]
        for fila in range(2, ws.max_row + 1):
            es_peso = ws.cell(fila, cols.index("Unidad") + 1).value == "$"
            for c in num_cols:
                ws.cell(fila, cols.index(c) + 1).number_format = '"$" #,##0' if es_peso else "#,##0.00"
            ws.cell(fila, cols.index("CV (%)") + 1).number_format = "0.0"
            ws.cell(fila, cols.index("Asimetría") + 1).number_format = "0.00"
            if ws.cell(fila, cols.index("|Asimetría| > 1") + 1).value == "sí":
                ws.cell(fila, cols.index("|Asimetría| > 1") + 1).fill = PatternFill("solid", fgColor="FCE8C8")
        for i, c in enumerate(cols, start=1):
            ws.cell(1, i).font = Font(bold=True)
            ws.cell(1, i).alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[get_column_letter(i)].width = 34 if c in ("Variable", "Recomendación") else 13
        ws.freeze_panes = "B2"
        wa = xw.sheets["Atraso"]
        wa.column_dimensions["A"].width = 44
        for fila in (3, 4):
            wa.cell(fila, 2).number_format = "0.0%"
    generados.append("descriptiva.xlsx")


# =============================================================================
# 3. Estadística inferencial (nivel familia)
# =============================================================================
def _cramer_v(t: np.ndarray) -> float:
    chi2 = stats.chi2_contingency(t, correction=False)[0]
    k = min(t.shape) - 1
    return math.sqrt(chi2 / (t.sum() * k)) if k else np.nan


def _chi2_monte_carlo(filas: pd.Series, cols: pd.Series, B: int, rng: np.random.Generator) -> tuple[float, float]:
    """p-valor por permutación de etiquetas con márgenes fijos (como chisq.test(simulate.p.value = TRUE) de R)."""
    fi, fl = pd.factorize(filas)
    ci, cl = pd.factorize(cols)
    r, c = len(fl), len(cl)

    def tabla(cc):
        return np.bincount(fi * c + cc, minlength=r * c).reshape(r, c)

    obs_t = tabla(ci)
    esperadas = np.outer(obs_t.sum(1), obs_t.sum(0)) / obs_t.sum()
    obs = ((obs_t - esperadas) ** 2 / esperadas).sum()
    extremos = 0
    cc = ci.copy()
    for _ in range(B):
        rng.shuffle(cc)
        if ((tabla(cc) - esperadas) ** 2 / esperadas).sum() >= obs - 1e-9:
            extremos += 1
    return float(obs), (1 + extremos) / (B + 1)


def _cruce_categorico(F: pd.DataFrame, var: str, etiqueta: str, orden: list, rng) -> dict:
    sub = F[F[var].notna()]
    t = pd.crosstab(sub[var], sub["moroso_lbl"]).reindex(index=[o for o in orden if o in set(sub[var])],
                                                          columns=[m for m in MOROSO if m in set(sub["moroso_lbl"])], fill_value=0)
    base = {"variable": etiqueta, "cruce": f"{etiqueta} × morosidad", "n": len(sub), "tabla": t, "puntaje": var == "riesgo_cat"}
    if t.shape[0] < 2 or t.shape[1] < 2:
        return {**base, "test": "—", "error": "una de las variables tiene una sola categoría con datos"}
    arr = t.to_numpy()
    chi2, p, gl, esp = stats.chi2_contingency(arr)  # corrección de Yates en 2x2
    V = _cramer_v(arr)
    if esp.min() >= 5:
        return {**base, "test": "Chi cuadrado (Yates)" if gl == 1 else "Chi cuadrado", "estadistico": chi2, "estadistico_tipo": "χ²",
                "gl": gl, "p": p, "efecto": V, "efecto_tipo": "V de Cramér"}
    if arr.shape == (2, 2):
        oddsr, p = stats.fisher_exact(arr)
        return {**base, "test": "Fisher exacto", "estadistico": oddsr, "estadistico_tipo": "odds ratio", "p": p,
                "efecto": V, "efecto_tipo": "V de Cramér"}
    chi2, p = _chi2_monte_carlo(sub[var], sub["moroso_lbl"], CONFIG["MONTE_CARLO"], rng)
    return {**base, "test": f"Chi cuadrado Monte Carlo ({CONFIG['MONTE_CARLO']} réplicas)", "estadistico": chi2,
            "estadistico_tipo": "χ²", "gl": gl, "p": p, "efecto": V, "efecto_tipo": "V de Cramér"}


def _spearman(F: pd.DataFrame, var: str, etiqueta: str) -> dict:
    sub = F[[var, "atraso_max"]].dropna()
    base = {"variable": etiqueta, "cruce": f"{etiqueta} vs. atraso máximo", "test": "Spearman", "n": len(sub),
            "puntaje": var == "puntaje_riesgo_actual"}
    if len(sub) < 4 or sub[var].nunique() < 2 or sub["atraso_max"].nunique() < 2:
        return {**base, "error": "menos de 4 pares o una variable constante"}
    rho, p = stats.spearmanr(sub[var], sub["atraso_max"])
    return {**base, "estadistico": rho, "estadistico_tipo": "rho", "p": p, "efecto": rho, "efecto_tipo": "rho"}


def _grupos(F: pd.DataFrame, var: str, etiqueta: str, orden: list) -> dict:
    presentes = [g for g in orden if (F[var] == g).any()]
    valores = [F.loc[F[var] == g, "atraso_max"].dropna().to_numpy() for g in presentes]
    medianas = [{"grupo": g, "n": len(v), "mediana": float(np.median(v)) if len(v) else np.nan} for g, v in zip(presentes, valores)]
    base = {"variable": etiqueta, "cruce": f"Atraso máximo según {etiqueta.lower()}", "n": int(sum(len(v) for v in valores)),
            "grupos": medianas, "puntaje": var == "riesgo_cat"}
    if len(presentes) < 2 or any(len(v) < 2 for v in valores):
        return {**base, "test": "—", "error": "se necesitan 2 o más grupos con 2 casos o más"}
    if len(presentes) == 2:
        a, b = valores
        u, p = stats.mannwhitneyu(a, b, alternative="two-sided", method="asymptotic", use_continuity=True)
        return {**base, "test": "Mann-Whitney", "estadistico": u, "estadistico_tipo": "U", "p": p,
                "efecto": 2 * u / (len(a) * len(b)) - 1, "efecto_tipo": "r rank-biserial"}
    if np.concatenate(valores).std() == 0:
        return {**base, "test": "Kruskal-Wallis", "error": "todos los valores son iguales"}
    h, p = stats.kruskal(*valores)
    n = sum(len(v) for v in valores)
    return {**base, "test": "Kruskal-Wallis", "estadistico": h, "estadistico_tipo": "H", "gl": len(presentes) - 1, "p": p,
            "efecto": h / (n - 1), "efecto_tipo": "ε²"}


def _auc(scores: np.ndarray, y: np.ndarray) -> float | None:
    pos, neg = y.sum(), len(y) - y.sum()
    if pos == 0 or neg == 0:
        return None
    r = stats.rankdata(scores)
    return float((r[y].sum() - pos * (pos + 1) / 2) / (pos * neg))


def _roc(F: pd.DataFrame, salida: Path, generados: list, omitidos: list) -> dict | None:
    scores = F["puntaje_riesgo_actual"].to_numpy(float)
    y = F["moroso"].to_numpy(bool)
    auc = _auc(scores, y)
    if auc is None:
        omitidos.append("roc.png: hace falta al menos una familia morosa y una no morosa")
        return None
    rng = np.random.default_rng(CONFIG["SEMILLA"] + 1)
    boot = []
    for _ in range(CONFIG["BOOTSTRAP"]):
        k = rng.integers(0, len(y), len(y))
        a = _auc(scores[k], y[k])
        if a is not None:
            boot.append(a)
    li, ls = np.percentile(boot, [2.5, 97.5])
    cortes = np.sort(np.unique(scores))[::-1]
    P, N = y.sum(), (~y).sum()
    fpr = [0.0] + [((scores >= c) & ~y).sum() / N for c in cortes]
    tpr = [0.0] + [((scores >= c) & y).sum() / P for c in cortes]

    fig, ax = plt.subplots(figsize=(5.6, 5))
    ax.plot([0, 1], [0, 1], ls="--", lw=1.2, color="#8A8F98", label="Azar (AUC 0,5)")
    ax.plot(fpr, tpr, lw=2, color="#1B4452", label=f"Puntaje de riesgo (AUC {fnum(auc, 3)})")
    ax.set_xlabel("1 − especificidad (falsos positivos)")
    ax.set_ylabel("Sensibilidad (verdaderos positivos)")
    ax.set_title(f"Curva ROC del puntaje de riesgo\nIC 95% bootstrap: {fnum(li, 3)} a {fnum(ls, 3)} · {P} morosas, {N} no morosas", fontsize=10)
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1.01)
    coma = FuncFormatter(lambda v, _: fnum(v, 1))
    ax.xaxis.set_major_formatter(coma)
    ax.yaxis.set_major_formatter(coma)
    ax.grid(alpha=0.3)
    ax.legend(loc="lower right", frameon=False)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    fig.savefig(salida / "roc.png", dpi=200)
    plt.close(fig)
    generados.append("roc.png")
    log.info("AUC del puntaje: %.3f (IC 95%%: %.3f a %.3f)", auc, li, ls)
    return {"auc": auc, "li": float(li), "ls": float(ls), "B": len(boot), "positivos": int(P), "negativos": int(N)}


def inferencial(F: pd.DataFrame, meta: dict, salida: Path, generados: list, omitidos: list) -> tuple[list, dict | None]:
    """Cruces contra la morosidad con test elegido según los datos, ajuste de Holm y ROC del puntaje."""
    rng = np.random.default_rng(CONFIG["SEMILLA"])
    ca, cb = CONFIG["CORTES_ASISTENCIA"]
    tests = [
        _cruce_categorico(F, "riesgo_cat", "Riesgo categorizado (P25/P75)", ORDEN_RIESGO, rng),
        _cruce_categorico(F, "beca_lbl", "Beca", ["Con beca", "Sin beca"], rng),
        _cruce_categorico(F, "asistencia_cat", "Asistencia categorizada", [f"Menos de {ca}%", f"{ca} a {cb}%", f"{cb}% o más"], rng),
        _spearman(F, "puntaje_riesgo_actual", "Puntaje de riesgo"),
        _spearman(F, "asistencia_prom", "Asistencia promedio"),
        _spearman(F, "cuota_mensual_familiar", "Cuota mensual familiar"),
        _grupos(F, "beca_lbl", "Beca", ["Con beca", "Sin beca"]),
        _grupos(F, "hijos_grupo", "Cantidad de hijos", list(dict.fromkeys(["1", "2", f"{meta['hijos_tope']} o más"]))),
        _grupos(F, "riesgo_cat", "Nivel de riesgo", ORDEN_RIESGO),
    ]
    idx = [i for i, t in enumerate(tests) if t.get("p") is not None and not pd.isna(t.get("p"))]
    if idx:
        ajust = multipletests([tests[i]["p"] for i in idx], method="holm")[1]
        for i, pa in zip(idx, ajust):
            tests[i]["p_ajustado"] = float(pa)
    for t in tests:
        t["significativo"] = t.get("p_ajustado") is not None and t["p_ajustado"] < CONFIG["ALFA"]
        if t.get("error"):
            log.warning("No se pudo correr %s: %s", t["cruce"], t["error"])
            omitidos.append(f"Test '{t['cruce']}': {t['error']}")
        else:
            log.info("%-48s %-28s p=%.4f p_aj=%.4f %s=%.3f", t["cruce"], t["test"], t["p"], t["p_ajustado"], t["efecto_tipo"], t["efecto"])

    tabla = pd.DataFrame([{
        "variable": t["cruce"], "test": t.get("test"), "estadistico": t.get("estadistico"), "tipo_estadistico": t.get("estadistico_tipo"),
        "gl": t.get("gl"), "p_valor": t.get("p"), "p_ajustado_holm": t.get("p_ajustado"), "tamano_efecto": t.get("efecto"),
        "tipo_efecto": t.get("efecto_tipo"), "n": t["n"], "significativo": "sí" if t["significativo"] else "no",
        "valida_el_puntaje": "sí" if t.get("puntaje") else "no", "nota": t.get("error", ""),
    } for t in tests])
    tabla.to_csv(salida / "inferencial.csv", sep=";", decimal=",", index=False, encoding="utf-8-sig", float_format="%.6g")
    generados.append("inferencial.csv")
    log.info("Asociación no implica causalidad; con n = %d familias los resultados son exploratorios.", len(F))
    return tests, _roc(F, salida, generados, omitidos)


# =============================================================================
# 4. Pareto (nivel familia)
# =============================================================================
def _lista_pareto(F: pd.DataFrame, valor: str, ordenar_por: str) -> pd.DataFrame:
    L = F.sort_values([ordenar_por, valor, "id_familia"], ascending=[False, False, True]).reset_index(drop=True)
    total = L[valor].sum()
    return pd.DataFrame({
        "rank": np.arange(1, len(L) + 1), "id_familia": L["id_familia"], "valor": L[valor],
        "pct_acumulado": L[valor].cumsum() / total if total else 0.0,
        "puntaje": L["puntaje_riesgo_actual"], "nivel_riesgo": L["riesgo_cat"],
    })


def _resumen_pareto(lista: pd.DataFrame) -> dict:
    total = lista["valor"].sum()
    ref = CONFIG["PARETO_REFERENCIA"]
    n_ref = int((lista["pct_acumulado"] < ref - 1e-12).sum() + 1) if total else None
    return {"total": float(total), "n": len(lista), "n_ref": n_ref,
            "pct_familias_ref": n_ref / len(lista) if n_ref else None,
            "pct_valor_ref": float(lista["pct_acumulado"].iloc[n_ref - 1]) if n_ref else None,
            "ceros": int((lista["valor"] == 0).sum()), "pct_ceros": float((lista["valor"] == 0).mean())}


def _grafico_pareto(lista: pd.DataFrame, r: dict, titulo: str, xlabel: str, ruta: Path) -> None:
    """Un solo eje en %: barra = participación de cada familia; línea = acumulado; referencia al 80%."""
    total = lista["valor"].sum()
    x = np.arange(1, len(lista) + 1)
    fig, ax = plt.subplots(figsize=(10, 4.6))
    ax.bar(x, lista["valor"] / total * 100, color="#1B4452", width=0.8, label="Participación de cada familia")
    ax.plot(x, lista["pct_acumulado"] * 100, color="#E9A23B", lw=2, marker="o", ms=3, label="Acumulado")
    ref = CONFIG["PARETO_REFERENCIA"] * 100
    ax.axhline(ref, ls="--", lw=1.2, color="#8A8F98", label=f"Referencia {fnum(ref, 0)}%")
    if r["n_ref"]:
        ax.axvline(r["n_ref"], ls=":", lw=1.2, color="#8A8F98")
        ax.annotate(f"{r['n_ref']} familias ({fpct(r['pct_familias_ref'])})\nexplican el {fpct(r['pct_valor_ref'])}",
                    xy=(r["n_ref"], ref), xytext=(r["n_ref"] + max(1, len(x) * 0.04), ref - 28), fontsize=9,
                    arrowprops={"arrowstyle": "-", "color": "#8A8F98"})
    ax.set_ylim(0, 102)
    ax.yaxis.set_major_formatter(PercentFormatter(decimals=0))
    ax.set_xlim(0.3, len(x) + 0.7)
    ax.set_xlabel(xlabel)
    ax.set_ylabel("% del total")
    ax.set_title(titulo)
    ax.grid(axis="y", alpha=0.3)
    ax.legend(loc="center right", frameon=False)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    fig.tight_layout()
    fig.savefig(ruta, dpi=200)
    plt.close(fig)


def pareto(F: pd.DataFrame, meta: dict, salida: Path, generados: list, omitidos: list) -> dict:
    """Atraso acumulado, deuda y concentración según el puntaje de riesgo."""
    out = {}
    analisis = [("atraso", "atraso_acum", "atraso_acum", "Pareto de los días de atraso acumulados por familia",
                 "Familias ordenadas por días de atraso acumulados", "del atraso")]
    if meta["hay_monto"]:
        analisis.append(("deuda", "monto_adeudado", "monto_adeudado", "Pareto de la deuda vencida por familia",
                         "Familias ordenadas por deuda", "de la deuda"))
    else:
        omitidos.append("pareto_deuda: no existe la columna monto_adeudado")
        log.warning("Sin columna monto_adeudado: se omite el Pareto de deuda")
    analisis.append(("puntaje", "atraso_acum", "puntaje_riesgo_actual", "Atraso acumulado con las familias ordenadas por puntaje de riesgo",
                     "Familias ordenadas por puntaje de riesgo (mayor a menor)", "del atraso"))
    for clave, valor, orden, titulo, xlabel, que in analisis:
        sub = F.copy()
        sub[valor] = sub[valor].fillna(0)
        lista = _lista_pareto(sub, valor, orden)
        r = _resumen_pareto(lista)
        if clave == "puntaje":
            top = max(1, math.ceil(len(lista) * CONFIG["TOP_PUNTAJE"]))
            ids_top = set(lista["id_familia"].iloc[:top])
            morosas = int(F["moroso"].sum())
            capt = int(F[F["moroso"] & F["id_familia"].isin(ids_top)].shape[0])
            r["top"] = {"n": top, "pct_familias": top / len(lista),
                        "pct_atraso": float(lista["pct_acumulado"].iloc[top - 1]) if r["total"] else None,
                        "morosas_capturadas": capt, "pct_morosas": capt / morosas if morosas else None}
            frase = (f"El {fpct(r['top']['pct_familias'], 0)} de familias con mayor puntaje ({top}) concentra el "
                     f"{fpct(r['top']['pct_atraso'])} del atraso y {capt} de las {morosas} familias morosas")
        else:
            uno = r["n_ref"] == 1
            frase = (f"{r['n_ref']} familia{'' if uno else 's'} ({fpct(r['pct_familias_ref'])} del total) explica{'' if uno else 'n'} el {fpct(r['pct_valor_ref'])} {que}"
                     if r["n_ref"] else f"No hay valores {que} para analizar")
        log.info("Pareto %s: %s. %d familias (%s) con valor cero.", clave, frase, r["ceros"], fpct(r["pct_ceros"]))
        r["frase"] = frase
        out[clave] = r
        lista.to_csv(salida / f"pareto_{clave}.csv", sep=";", decimal=",", index=False, encoding="utf-8-sig", float_format="%.4f")
        generados.append(f"pareto_{clave}.csv")
        if r["total"]:
            _grafico_pareto(lista, r, titulo, xlabel, salida / f"pareto_{clave}.png")
            generados.append(f"pareto_{clave}.png")
        else:
            omitidos.append(f"pareto_{clave}.png: el total es cero, no hay nada que concentrar")
    return out


# =============================================================================
# 5. Conclusiones automáticas
# =============================================================================
def conclusiones(fuente: str, valid: pd.DataFrame, desc: pd.DataFrame, extra: dict, F: pd.DataFrame, meta: dict,
                 tests: list, roc: dict | None, par: dict, salida: Path, generados: list) -> None:
    """Escribe conclusiones.md sólo con números calculados en este script."""
    umbral, alfa = CONFIG["UMBRAL_MOROSO_DIAS"], CONFIG["ALFA"]
    L = ["# Conclusiones del análisis de morosidad", "",
         f"Fuente: {fuente}. Familia morosa: al menos una cuota con más de {umbral} días de atraso. "
         f"Nivel de significación: {fnum(alfa, 2)} (p ajustado por Holm).", ""]
    prob = valid[~valid["ok"]]
    L += ["## Calidad de los datos", ""]
    L += [f"- {r.tabla}: {r.chequeo.lower()} → {r.casos}{f' ({r.detalle})' if r.detalle else ''}." for r in prob.itertuples()] \
        if len(prob) else ["- Sin nulos, duplicados ni valores fuera de rango."]
    L.append("")

    def var(nombre):
        f = desc[desc["variable"] == nombre]
        return f.iloc[0] if len(f) else None

    L += ["## Hallazgos descriptivos", ""]
    at, pu, cu, asis = var("Días de atraso por cuota"), var("Puntaje de riesgo actual"), var("Cuota mensual familiar"), var("Asistencia últimos 30 días")
    nf_ = var("Hijos por familia")
    na = var("Edad")
    L.append(f"- Se analizaron {extra['n']} cuotas, {int(nf_['N']) if nf_ is not None else '—'} familias y {int(na['N']) if na is not None else '—'} alumnos.")
    if at is not None:
        L.append(f"- El {fpct(extra['pct_mayor_0'])} de las cuotas se pagó (o sigue impaga) con algún día de atraso y el "
                 f"{fpct(extra['pct_mayor_umbral'])} superó los {umbral} días. Mediana del atraso: {fnum(at['mediana'])} días; "
                 f"P75: {fnum(at['P75'])}; máximo: {fnum(at['max'], 0)}.")
    if pu is not None and pu["N"]:
        L.append(f"- Puntaje de riesgo: mediana {fnum(pu['mediana'])}, P25 {fnum(pu['P25'])}, P75 {fnum(pu['P75'])} (media {fnum(pu['media'])}).")
    if cu is not None and cu["N"]:
        L.append(f"- Cuota mensual familiar: mediana {fmoney(cu['mediana'])}, entre {fmoney(cu['min'])} y {fmoney(cu['max'])}.")
    if asis is not None and asis["N"]:
        L.append(f"- Asistencia de los alumnos en 30 días: mediana {fnum(asis['mediana'])}%, P25 {fnum(asis['P25'])}%.")
    asim = desc[desc["asimetria_marcada"] == "sí"]
    if len(asim):
        L.append("- Variables con asimetría marcada (|asimetría| > 1), donde conviene informar mediana y percentiles en lugar de la media: "
                 + ", ".join(f"{r.variable.lower()} ({fnum(r.asimetria, 2)})" for r in asim.itertuples()) + ".")
    n_mor = int(F["moroso"].sum())
    L += [f"- {n_mor} de {len(F)} familias ({fpct(n_mor / len(F) if len(F) else None)}) quedaron clasificadas como morosas.", ""]

    L += ["## Asociaciones (nivel familia)", ""]

    def linea(t):
        return (f"- {t['cruce']} — {t['test']}, {fp('p', t['p'])}, {fp('p ajustado', t['p_ajustado'])}, "
                f"{t['efecto_tipo']} = {fnum(t['efecto'], 2)}, n = {t['n']}.")

    sig = [t for t in tests if t["significativo"]]
    nosig = [t for t in tests if not t["significativo"] and t.get("p") is not None and not t.get("error")]
    sin = [t for t in tests if t.get("error")]
    L += ["**Con asociación significativa después de ajustar por comparaciones múltiples:**", ""]
    L += [linea(t) for t in sig] or ["- Ninguna."]
    L += ["", "**Sin asociación significativa:**", ""]
    L += [linea(t) for t in nosig] or ["- Ninguna."]
    L.append("")
    if sin:
        L += ["**No se pudieron correr:**", ""] + [f"- {t['cruce']}: {t['error']}." for t in sin] + [""]
    L += ["> Asociación no implica causalidad. Con este tamaño de muestra los resultados son exploratorios: sirven para orientar, no para confirmar.", ""]

    L += ["## Capacidad discriminante del puntaje de riesgo", ""]
    if roc:
        a = roc["auc"]
        lect = "excelente" if a >= 0.9 else "buena" if a >= 0.8 else "aceptable" if a >= 0.7 else "pobre" if a >= 0.6 else "casi nula"
        L += [f"- AUC = {fnum(a, 3)} (IC 95% bootstrap: {fnum(roc['li'], 3)} a {fnum(roc['ls'], 3)}; {roc['B']} réplicas), con "
              f"{roc['positivos']} familias morosas y {roc['negativos']} no morosas: discriminación {lect}.",
              "- Advertencia: el puntaje actual incorpora el historial de pagos y la deuda vencida, que son parte de lo que define a una "
              "familia como morosa. Este AUC es optimista; la validación honesta usa el puntaje calculado antes de cada vencimiento.", ""]
    else:
        L += ["- No se pudo calcular: hace falta al menos una familia morosa y una no morosa.", ""]

    L += ["## Concentración (Pareto)", ""]
    for clave in ("atraso", "deuda"):
        if clave in par:
            r = par[clave]
            L.append(f"- {r['frase']}. {r['ceros']} familias ({fpct(r['pct_ceros'])}) tienen valor cero.")
        elif clave == "deuda":
            L.append("- Pareto de deuda omitido: no hay columna de monto adeudado.")
    t = par["puntaje"]["top"]
    L += [f"- El {fpct(t['pct_familias'], 0)} de familias con mayor puntaje ({t['n']}) concentra el {fpct(t['pct_atraso'])} de los días de "
          f"atraso y {t['morosas_capturadas']} de las {n_mor} familias morosas ({fpct(t['pct_morosas'])}).", ""]

    L += ["## Implicancias para el diseño del sistema", ""]
    if par["atraso"]["n_ref"]:
        L.append(f"- **Focalización:** la gestión intensiva puede concentrarse en {par['atraso']['n_ref']} familias; el resto se atiende con recordatorios automáticos.")
    u = CONFIG["UMBRALES_SEMAFORO_ACTUALES"]
    L.append(f"- **Cortes del semáforo:** en la distribución real, P25 = {fnum(meta['p25'])} y P75 = {fnum(meta['p75'])} puntos"
             + (f"; hoy el sistema usa medio desde {u['medio']} y alto desde {u['alto']}" if u else "")
             + ". Conviene revisar los cortes para que \"alto\" agrupe a la cola real de riesgo, y recalibrarlos cada ciclo.")
    del_p = [x for x in sig if x.get("puntaje")]
    externas = list(dict.fromkeys(x["variable"].lower() for x in sig if not x.get("puntaje")))
    if del_p:
        L.append(f"- **Puntaje como alerta:** el puntaje ordena a las familias según su atraso ({len(del_p)} de "
                 f"{sum(1 for x in tests if x.get('puntaje'))} cruces significativos); sigue siendo la señal principal, con la salvedad de circularidad indicada arriba.")
    L.append("- **Otras señales de alerta temprana:** " + (f"con respaldo en estos datos: {', '.join(externas)}." if externas else
             "ninguna variable fuera del puntaje (beca, asistencia, cuota, cantidad de hijos) pasó el ajuste por comparaciones múltiples; "
             "no conviene sumarlas como alerta hasta tener más datos."))
    if any("asistencia" in x["variable"].lower() for x in sig):
        L.append("- La asistencia no se usa en el puntaje de cobranza por principio; si se asocia con la morosidad, es una señal para que "
                 "la dirección acompañe a la familia, nunca para cobrar.")
    L.append("")

    L += ["## Limitaciones", ""]
    L.append(f"- Muestra chica: {len(F)} familias analizadas ({n_mor} morosas) y {int(F['beca'].sum())} familias con algún hijo becado; "
             "los tests tienen poca potencia para detectar efectos moderados.")
    if meta["familias_total"] > len(F):
        L.append(f"- {meta['familias_total'] - len(F)} familias sin cuotas vencidas (o sin puntaje) quedaron fuera de los análisis a nivel familia.")
    L.append("- Foto estática: se cruzan variables actuales con el atraso observado; no hay dimensión temporal ni se controla por otras variables.")
    if fuente == "sintético":
        L.append("- **Datos sintéticos:** estos resultados sólo sirven para probar el pipeline; no describen ninguna escuela real.")
    (salida / "conclusiones.md").write_text("\n".join(L) + "\n", encoding="utf-8")
    generados.append("conclusiones.md")


# =============================================================================
# main
# =============================================================================
def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Análisis estadístico de morosidad (descriptiva, inferencial, Pareto y conclusiones)")
    ap.add_argument("--alumnos")
    ap.add_argument("--familias")
    ap.add_argument("--cuotas")
    ap.add_argument("--archivo", help="un único CSV/Excel (tres hojas o una tabla plana)")
    ap.add_argument("--sintetico", action="store_true", help="usar el dataset sintético reproducible")
    ap.add_argument("--salida", default=CONFIG["SALIDA"])
    args = ap.parse_args(argv)
    for flujo in (sys.stdout, sys.stderr):  # consolas de Windows en cp1252: que los acentos no rompan la salida
        try:
            flujo.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    logging.basicConfig(level=logging.INFO, format="%(levelname)-7s %(message)s")
    np.random.seed(CONFIG["SEMILLA"])

    salida = Path(args.salida)
    salida.mkdir(parents=True, exist_ok=True)
    generados: list[str] = []
    omitidos: list[str] = []

    datos, fuente = cargar_datos(args)
    valid = validar(datos)
    valid.to_csv(salida / "validacion.csv", sep=";", index=False, encoding="utf-8-sig")
    generados.append("validacion.csv")
    datos = limpiar(datos)
    desc, extra = descriptiva(datos, salida, generados, omitidos)
    F, meta = tabla_familias(datos)
    if len(F) < 4:
        log.error("Sólo %d familias con cuotas vencidas: no alcanza para los análisis a nivel familia", len(F))
        omitidos.append("inferencial, pareto y conclusiones: menos de 4 familias con cuotas vencidas")
    else:
        tests, roc = inferencial(F, meta, salida, generados, omitidos)
        par = pareto(F, meta, salida, generados, omitidos)
        conclusiones(fuente, valid, desc, extra, F, meta, tests, roc, par, salida, generados)

    print("\n" + "=" * 70)
    print(f"Resumen · fuente de datos: {fuente}")
    print(f"Carpeta de salida: {salida.resolve()}")
    print("Generado:")
    for g in generados:
        print(f"  + {g}")
    print("Omitido:" if omitidos else "Omitido: nada")
    for o in omitidos:
        print(f"  - {o}")
    print("=" * 70)
    return 0


if __name__ == "__main__":
    sys.exit(main())
